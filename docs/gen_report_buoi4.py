# -*- coding: utf-8 -*-
"""
Sinh bao cao thuc hanh Buoi 4 - Deployment & Testing (.docx) cho du an
Smart City Bus Assistant.

Chay: python docs/gen_report_buoi4.py
Ket qua: docs/BaoCao_Buoi4_Deployment_Testing.docx
"""
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "BaoCao_Buoi4_Deployment_Testing.docx")

ACCENT = RGBColor(0x1E, 0x40, 0xAF)   # xanh dam
HEADER_BG = "1E40AF"

doc = Document()

# --- Font mac dinh ---
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def h1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = ACCENT
    p.space_before = Pt(12)
    p.space_after = Pt(4)
    return p


def h2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = ACCENT
    return p


def para(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def make_table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)
        shade_cell(hdr[i], HEADER_BG)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = w
    doc.add_paragraph()
    return t


# ======================================================================
# TRANG BIA
# ======================================================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("TRƯỜNG ĐẠI HỌC LẠC HỒNG")
r.bold = True
r.font.size = Pt(13)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("KHOA CÔNG NGHỆ THÔNG TIN").font.size = Pt(11)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BÁO CÁO THỰC HÀNH BUỔI 4")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = ACCENT
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("CHỦ ĐỀ: DEPLOYMENT & TESTING (TRIỂN KHAI CLOUD & KIỂM THỬ)")
r.bold = True
r.font.size = Pt(12)

doc.add_paragraph()
info = doc.add_table(rows=4, cols=2)
info.style = "Table Grid"
info_data = [
    ("Đề tài:", "Smart City Bus Assistant – Cổng tra cứu tuyến & theo dõi xe buýt đô thị "
                "(TP.HCM – Biên Hòa – Kiên Giang)"),
    ("Sinh viên thực hiện:", "[Điền họ và tên]"),
    ("Mã số sinh viên / Lớp:", "[Điền MSSV] / [Điền lớp]"),
    ("Ngành học:", "[Điền ngành]"),
]
for i, (k, v) in enumerate(info_data):
    c0, c1 = info.rows[i].cells
    c0.text = ""
    c0.paragraphs[0].add_run(k).bold = True
    c1.text = v

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Năm học 2025 – 2026").italic = True

doc.add_page_break()

# ======================================================================
# 1. TONG QUAN
# ======================================================================
h1("1. TỔNG QUAN HỆ THỐNG & ĐỊA CHỈ TRUY CẬP CLOUD (OUTPUT)")
para(
    "Smart City Bus Assistant là ứng dụng điện toán đám mây giúp người dùng (sinh viên, người "
    "cao tuổi, người đi làm bằng phương tiện công cộng) tra cứu tuyến xe buýt tối ưu giữa hai "
    "điểm, ước tính thời gian – chi phí và theo dõi vị trí xe (mô phỏng) trên bản đồ. Hệ thống "
    "được triển khai 100% trên nền tảng điện toán đám mây với 4 thành phần Cloud có vai trò thực "
    "tế: Cloud Hosting, Cloud Database, Cloud API và Cloud Storage."
)
para(
    "Khác với kiến trúc tách rời Frontend/Backend chạy trên 2 nền tảng riêng, ứng dụng dùng "
    "Streamlit – một framework Python hợp nhất giao diện và xử lý nghiệp vụ trong cùng một tiến "
    "trình; phần logic nghiệp vụ (thuật toán tìm tuyến, truy xuất dữ liệu, mô phỏng vị trí xe) "
    "được tách riêng thành package backend/ theo mô hình Monorepo phân tầng."
)
make_table(
    ["Hạng mục (Output)", "Nền tảng / Dịch vụ Cloud", "Đường dẫn / Thông tin truy cập"],
    [
        ["Web Application (Frontend + Backend)", "Streamlit Community Cloud",
         "https://busroutecloud.streamlit.app/"],
        ["Cloud Database", "Supabase – Managed PostgreSQL 15 (Region: ap-southeast-1, Singapore)",
         "Project ref: pieplfirsanegkuxojch\nHost: db.pieplfirsanegkuxojch.supabase.co"],
        ["Cloud API", "Supabase PostgREST (REST API tự sinh trên mọi bảng)",
         "https://pieplfirsanegkuxojch.supabase.co/rest/v1/"],
        ["Cloud Storage", "Supabase Storage – bucket công khai \"datasets\"",
         "https://pieplfirsanegkuxojch.supabase.co/storage/v1/object/public/datasets/"],
        ["Source Code Repository", "GitHub (Monorepo)",
         "https://github.com/Duongtruongtai/busroute_cloud"],
        ["Tài khoản đăng nhập", "Không áp dụng",
         "App tra cứu công cộng – không yêu cầu đăng nhập (giống Google Maps / BusMap)"],
    ],
)

# ======================================================================
# 2. DEPLOYMENT
# ======================================================================
h1("2. BÁO CÁO CHI TIẾT TRIỂN KHAI ĐÁM MÂY (DEPLOYMENT)")
para(
    "Thực hiện theo yêu cầu chuẩn hóa các thành phần điện toán đám mây, nhóm đã hoàn thiện quá "
    "trình cấu hình và tự động hóa build/deploy (CI/CD) như sau:"
)

h2("2.1. Web Application → Cloud Hosting (Streamlit Community Cloud)")
bullet("Công nghệ: Python 3.11, Streamlit, pandas, folium + streamlit-folium (bản đồ), "
       "streamlit-searchbox (gợi ý địa điểm), streamlit-autorefresh (mô phỏng realtime), "
       "supabase-py (kết nối Cloud Database).")
bullet("Cấu trúc Monorepo: main file frontend/app.py; logic nghiệp vụ tách riêng ở backend/ "
       "(route_finder.py – thuật toán tìm tuyến, datastore.py – truy xuất dữ liệu, schedule.py – "
       "giờ hoạt động, tracking.py – mô phỏng vị trí xe, search.py & geocoding.py – tìm kiếm, "
       "i18n.py – đa ngôn ngữ).")
bullet("CI/CD tự động: Streamlit Cloud kết nối trực tiếp GitHub, tự động cài đặt requirements.txt "
       "và redeploy mỗi khi có commit mới trên branch master – không cần thao tác thủ công.")
bullet("Biến môi trường (Secrets): SUPABASE_URL và SUPABASE_KEY (anon public key) được cấu hình "
       "trong mục App settings → Secrets của Streamlit Cloud, KHÔNG commit lên GitHub "
       "(.streamlit/secrets.toml nằm trong .gitignore).")
bullet("Đặc tính Cold-start: gói Community Cloud miễn phí sẽ đưa app vào trạng thái ngủ sau một "
       "thời gian dài không có truy cập; lần truy cập đầu tiên sau đó mất ~30 giây để khởi động "
       "lại. Đã ghi nhận là hạn chế và đề xuất khắc phục ở mục 6.")

h2("2.2. Business Logic & Lớp truy xuất dữ liệu")
bullet("Thuật toán tìm tuyến (backend/route_finder.py): tìm tuyến TRỰC TIẾP trước, sau đó tìm "
       "phương án chuyển tuyến TỐI ĐA 1 LẦN thông qua các trạm trung chuyển dùng chung (bến xe). "
       "Tuyến được mô hình hóa CHẠY 2 CHIỀU đúng như xe buýt thực tế. Xếp hạng kết quả theo số "
       "lần chuyển tuyến rồi đến tổng thời gian; gắn nhãn phương án Nhanh nhất / Rẻ nhất / Ít "
       "chuyển tuyến nhất.")
bullet("Kết nối dữ liệu (backend/datastore.py): gọi Cloud Database qua supabase-py; nếu Cloud "
       "tạm mất kết nối, tự động chuyển sang đọc file CSV cục bộ (cơ chế fallback) để ứng dụng "
       "không bị gián đoạn.")
bullet("Ghi dữ liệu lên Cloud: mỗi lượt tra cứu được INSERT một dòng vào bảng search_logs "
       "(điểm đi, điểm đến, loại vé, số kết quả, thời điểm) – chứng minh ứng dụng GHI dữ liệu "
       "thật lên Cloud Database chứ không chỉ đọc.")
bullet("Mô phỏng vị trí xe (backend/tracking.py): tính vị trí ước tính của từng chuyến xe đang "
       "chạy dựa trên giờ khởi hành chuẩn + giãn cách + thời gian đã trôi qua, nội suy theo lộ "
       "trình. LUÔN được ghi rõ trên giao diện là ƯỚC TÍNH MÔ PHỎNG, không phải GPS thời gian "
       "thực (Việt Nam hiện chưa có API GPS xe buýt công khai miễn phí).")

h2("2.3. Database → Cloud Database (Supabase PostgreSQL)")
bullet("Công nghệ: PostgreSQL 15 quản trị bởi Supabase Cloud, Region Singapore (ap-southeast-1) "
       "– gần Việt Nam nên độ trễ thấp.")
bullet("Kết nối: ứng dụng truy cập database qua REST API (PostgREST) bằng thư viện supabase-py "
       "với anon public key (không kết nối trực tiếp cổng PostgreSQL từ client).")
bullet("Bảo mật Row Level Security (RLS): bật RLS trên cả 4 bảng. Chính sách: cho phép ĐỌC công "
       "khai 4 bảng (dữ liệu giao thông công cộng, không nhạy cảm); chỉ cho phép GHI công khai "
       "riêng bảng search_logs. Các thao tác sửa/xóa dữ liệu tuyến chỉ thực hiện được bằng "
       "service_role key phía máy chủ.")
bullet("Schema & Migration: thực thi bộ SQL (database/schema.sql) tạo 4 bảng: stops, routes, "
       "route_stops, search_logs, kèm index và policy RLS. Script an toàn khi chạy lại nhiều lần "
       "(create table if not exists + alter table add column if not exists).")
bullet("Seed Data: 426 trạm, 24 tuyến, 455 dòng route_stops đã được nạp sẵn phục vụ nghiệm thu "
       "(script database/seed_supabase.py).")

h2("2.4. Storage → Cloud Storage (Supabase Storage)")
bullet("Giải pháp: bucket công khai \"datasets\" trên Supabase Storage, lưu bản sao lưu 3 file "
       "dữ liệu gốc: stops.csv, routes.csv, route_stops.csv.")
bullet("Vai trò thực tế: sao lưu / versioning bộ dataset; cho phép tải lại dữ liệu gốc bất kỳ "
       "lúc nào qua đường dẫn HTTPS công khai (phân phối qua CDN của Supabase). "
       "Ví dụ: .../storage/v1/object/public/datasets/stops.csv")
bullet("Bảo mật: việc upload/ghi đè file chỉ thực hiện bằng service_role key phía máy chủ khi "
       "seed dữ liệu; key này KHÔNG commit lên GitHub và KHÔNG có trong bản build phía client.")

h2("2.5. Dịch vụ bên thứ ba miễn phí được tích hợp")
bullet("OpenStreetMap tile server: cung cấp nền bản đồ (đường, sông ngòi, nhãn địa danh) – "
       "miễn phí, không cần API key.")
bullet("Nominatim (OpenStreetMap): dịch vụ geocoding chuyển địa chỉ tự do thành tọa độ, dùng "
       "làm PHƯƠNG ÁN DỰ PHÒNG khi tên người dùng gõ không khớp trực tiếp tên trạm trong dữ liệu "
       "(có giới hạn tần suất gọi để tôn trọng chính sách sử dụng).")

# ======================================================================
# 3. TEST CASES
# ======================================================================
h1("3. KỊCH BẢN & KẾT QUẢ KIỂM THỬ HỆ THỐNG (TEST CASES & RESULTS)")
para(
    "Nhóm xây dựng bộ 12 Test Case bao phủ các luồng chức năng quan trọng: đọc/ghi Cloud "
    "Database, thuật toán tìm tuyến (trực tiếp, chuyển tuyến, 2 chiều), gợi ý địa điểm, lọc theo "
    "khu vực, hiển thị bản đồ, mô phỏng vị trí xe, xử lý lỗi và truy cập Cloud Storage. "
    "Các test tự động cho thuật toán được kiểm chứng bằng file tests/test_route_finder.py "
    "(chạy: python tests/test_route_finder.py – kết quả PASS toàn bộ)."
)
make_table(
    ["Mã", "Mục tiêu kiểm thử", "Các bước thực hiện", "Kết quả mong đợi", "Thực tế", "Đánh giá"],
    [
        ["TC-01", "Đọc dữ liệu tuyến/trạm từ Cloud Database",
         "Mở app, quan sát số tuyến & số trạm ở thanh bên",
         "App gọi Supabase, hiển thị đúng 24 tuyến / 426 trạm",
         "Badge dữ liệu tải đầy đủ từ Supabase", "PASS"],
        ["TC-02", "Ghi nhật ký tìm kiếm lên Cloud Database",
         "Thực hiện 1 lượt tra cứu bất kỳ",
         "Thêm 1 dòng mới vào bảng search_logs trên Supabase",
         "Số dòng search_logs tăng sau mỗi lượt tra cứu", "PASS"],
        ["TC-03", "Tìm tuyến TRỰC TIẾP",
         "Điểm đi: Chợ Bến Thành → Điểm đến: Bến xe Chợ Lớn, bấm Tìm tuyến",
         "Trả về phương án 0 lần chuyển tuyến, giá vé sinh viên 3.000đ",
         "Ra tuyến 01, 0 lần chuyển, đúng giá vé", "PASS"],
        ["TC-04", "Tìm tuyến CÓ CHUYỂN TUYẾN",
         "Điểm đi: Bến xe Miền Tây → Điểm đến: Đại học Quốc Gia",
         "Trả về phương án 1 lần chuyển, điểm chuyển là Bến Thành",
         "Đúng 1 lần chuyển tại Bến Thành", "PASS"],
        ["TC-05", "Tuyến chạy 2 CHIỀU (khứ hồi)",
         "Điểm đi: Đại học Bách Khoa → Điểm đến: Bến Thành (ngược chiều mô tả tuyến)",
         "Vẫn tìm được phương án (xe buýt thực tế chạy 2 chiều)",
         "Ra 3 phương án hợp lệ", "PASS"],
        ["TC-06", "Kịch bản demo dữ liệu THẬT (Kiên Giang)",
         "Chọn khu vực Kiên Giang, Điểm đi: Bưu điện Kiên Lương → Bến xe Hòn Đất",
         "Ra phương án chuyển tuyến tại Bến xe Tri Tôn, ~44 phút",
         "Đúng như kỳ vọng, chuyển tuyến tại Tri Tôn", "PASS"],
        ["TC-07", "Gợi ý địa điểm khi gõ (autocomplete)",
         "Gõ \"nguyen tri phuong\" vào ô Điểm đi",
         "Hiện gợi ý \"Trường Nguyễn Tri Phương (đường Nguyễn Ái Quốc)\"",
         "Gợi ý xuất hiện, không phân biệt dấu", "PASS"],
        ["TC-08", "Gõ có từ thừa vẫn ra gợi ý",
         "Gõ \"tram xe buyt Kien Luong\" (có từ không có trong tên trạm)",
         "Vẫn hiện các trạm gần khớp thay vì \"No options\"",
         "Trả về các trạm chứa \"Kiên Lương\"", "PASS"],
        ["TC-09", "Lọc theo khu vực",
         "Bấm nút khu vực \"Biên Hòa - Đồng Nai\"",
         "Chỉ còn hiển thị 5 tuyến / dữ liệu Biên Hòa; gợi ý cũng giới hạn theo khu vực",
         "Danh sách tuyến/trạm và metric cập nhật đúng", "PASS"],
        ["TC-10", "Bản đồ & hiệu ứng marker điểm đi/đến",
         "Sau khi tìm tuyến, quan sát bản đồ",
         "Bản đồ zoom sát 2 điểm, hiện 2 marker nhấp nháy (pulse), không có marker rác",
         "Đúng: chỉ 2 marker pulse, zoom vừa khít", "PASS"],
        ["TC-11", "Mô phỏng vị trí xe (simulated realtime)",
         "Bật \"Tự động cập nhật vị trí xe\", chọn tuyến đang trong giờ hoạt động",
         "Bản đồ hiện icon xe di chuyển, tự làm mới; có cảnh báo \"ước tính mô phỏng\"",
         "Icon xe hiển thị, làm mới mỗi 8s, có ghi chú rõ", "PASS"],
        ["TC-12", "Truy cập Cloud Storage (bản sao lưu dataset)",
         "Mở URL công khai .../datasets/stops.csv",
         "Trả về file CSV qua HTTPS, mã 200",
         "HTTP 200, tải về 43.588 bytes", "PASS"],
    ],
)
para("Ngoài ra, các trường hợp biên đã được kiểm tra và xử lý mềm mại (không hiện traceback): "
     "điểm đi trùng điểm đến → báo lỗi thân thiện; không tìm được tuyến → hiện gợi ý thử lại; "
     "mất kết nối Cloud → tự chuyển sang dữ liệu cục bộ.")

# ======================================================================
# 4. EVALUATION MATRIX
# ======================================================================
h1("4. ĐÁNH GIÁ TOÀN DIỆN NĂNG LỰC HỆ THỐNG (EVALUATION MATRIX)")

h2("4.1. Khả năng mở rộng (Scalability)")
bullet("Web App: Streamlit Community Cloud là nền tảng được quản lý (managed); khi cần nhiều tài "
       "nguyên hơn có thể chuyển sang tự host trên VPS/Container hoặc Streamlit trong Snowflake. "
       "Hạn chế hiện tại: gói miễn phí chạy 1 instance, chưa auto-scale ngang.")
bullet("Database: Supabase tích hợp Connection Pooling (PgBouncer) xử lý tốt hàng nghìn kết nối "
       "đồng thời; có thể bật Read Replicas và nâng dung lượng khi mở rộng quy mô dữ liệu.")
bullet("Dữ liệu: kiến trúc dựa trên bảng route_stops (quan hệ nhiều-nhiều tuyến–trạm) cho phép "
       "thêm tỉnh/thành mới chỉ bằng cách nạp thêm dữ liệu, không phải sửa mã nguồn.")

h2("4.2. Độ ổn định & Tính sẵn sàng (Availability)")
bullet("Supabase cam kết uptime cao (~99.9%) cho hạ tầng database và storage.")
bullet("Điểm lưu ý: Streamlit Community Cloud có cơ chế \"ngủ\" sau thời gian dài không truy cập "
       "(tương tự Render Free Tier). Lần truy cập đầu tiên sau khi ngủ mất ~30 giây để khởi động "
       "lại. Khi tự host hoặc dùng gói trả phí, app hoạt động liên tục 24/7.")
bullet("Cơ chế fallback: nếu Cloud Database gặp sự cố, ứng dụng tự chuyển sang đọc dữ liệu CSV "
       "cục bộ (đã đóng gói trong bản build) nên vẫn tra cứu được.")

h2("4.3. Hiệu năng & Thời gian phản hồi (Performance) – số liệu đo thực tế")
make_table(
    ["Hạng mục đo", "Kết quả", "Ghi chú"],
    [
        ["Tải khung trang Web App", "~2,4 – 2,6 giây",
         "Gồm 3 bước bắt tay của Streamlit Cloud + tải HTML shell"],
        ["Truy vấn Supabase – bảng routes / route_stops / search_logs", "~90 – 105 ms/truy vấn",
         "Đo qua supabase-py, 3 lần lấy trung bình"],
        ["Tải toàn bộ 426 trạm từ Supabase", "~136 ms",
         "Thao tác app thực hiện khi khởi động (có cache 5 phút)"],
        ["Thuật toán tìm tuyến (route_finder)", "< 50 ms",
         "Dữ liệu ~426 trạm / 24 tuyến, tìm trực tiếp + 1 lần chuyển"],
        ["Tải file từ Cloud Storage (stops.csv)", "~0,66 giây",
         "Qua HTTPS/CDN công khai của Supabase"],
        ["Nền bản đồ (OpenStreetMap) + geocoding (Nominatim)", "~0,5 – 1 giây",
         "Tùy chất lượng mạng; geocoding chỉ gọi khi cần dự phòng"],
    ],
)

h2("4.4. An toàn & Bảo mật (Security)")
bullet("Mã hóa đường truyền: 100% kết nối qua HTTPS / TLS – cả domain Streamlit Cloud, Supabase "
       "REST API và Supabase Storage.")
bullet("Phân quyền dữ liệu: bật Row Level Security trên toàn bộ bảng. Anon public key (dùng phía "
       "client) chỉ có quyền đọc 4 bảng và ghi riêng bảng search_logs; không thể sửa/xóa dữ liệu "
       "tuyến.")
bullet("Quản lý bí mật: service_role key (toàn quyền) chỉ dùng phía máy chủ khi seed/upload dữ "
       "liệu, hoàn toàn không commit lên GitHub và không nằm trong bản build client. "
       "File .streamlit/secrets.toml nằm trong .gitignore.")
bullet("Không lưu thông tin cá nhân (PII): bảng search_logs chỉ ghi điểm đi/đến và thời điểm tra "
       "cứu – không lưu địa chỉ IP, không lưu tài khoản (ứng dụng không có đăng nhập).")

h2("4.5. Dự toán chi phí triển khai (Cost Estimation)")
make_table(
    ["Dịch vụ Cloud", "Giai đoạn MVP (hiện tại)", "Giai đoạn sản xuất (Production)", "Ghi chú"],
    [
        ["Web App (Streamlit Community Cloud)", "0 đ (Free)", "~120.000 – 150.000 đ/tháng",
         "Bản Free đủ cho demo; production nên tự host VPS ~5–6 USD/tháng để chạy 24/7"],
        ["Cloud Database (Supabase)", "0 đ (Free – 500MB DB)", "~625.000 đ/tháng (Pro – 25 USD)",
         "Bản Free: 500MB database, 1GB storage, đủ cho quy mô đồ án"],
        ["Cloud Storage (Supabase Storage)", "0 đ (nằm trong gói Free)", "Nằm trong gói Supabase Pro",
         "Bản sao lưu dataset dung lượng nhỏ (~60 KB)"],
        ["Cloud API (Supabase PostgREST)", "0 đ (đi kèm database)", "Đi kèm gói Supabase",
         "Không tính phí riêng"],
        ["OpenStreetMap / Nominatim", "0 đ", "0 đ",
         "Miễn phí theo chính sách sử dụng (usage policy)"],
        ["TỔNG CHI PHÍ", "0 đ (hoàn toàn miễn phí)", "~750.000 – 780.000 đ/tháng (~30 USD)",
         "Tối ưu chi phí tối đa cho quy mô sinh viên"],
    ],
)

h2("4.6. Khả năng tương thích & Đa thiết bị (Accessibility)")
bullet("Giao diện Streamlit tự động responsive (layout wide + các cột co giãn); test hiển thị "
       "tốt trên trình duyệt máy tính (Chrome, Edge, Firefox, Cốc Cốc) và trên điện thoại.")
bullet("Song ngữ Việt / Anh (chuyển ngay trên thanh bên) phục vụ cả người dùng nước ngoài.")
bullet("Chế độ giao diện Sáng / Tối; bản đồ tự đổi nền theo giao diện.")
bullet("Toàn bộ chữ tiếng Việt hiển thị đầy đủ dấu.")

# ======================================================================
# 5. DATASET
# ======================================================================
h1("5. BỘ DỮ LIỆU ĐÃ SEED VÀO DATABASE (DATASET)")
para(
    "Cơ sở dữ liệu Supabase được khởi tạo với bộ dữ liệu phục vụ nghiệm thu tức thì. Điểm nổi "
    "bật: dữ liệu Kiên Giang là DỮ LIỆU THẬT (không phải tự tạo) – trích xuất từ file trạm xe "
    "buýt công khai của tỉnh, đã qua xử lý (loại 4 lần trùng lặp trong nguồn, tách thứ tự trạm "
    "từ số Km / số điểm dừng ghi trong tên, định vị tọa độ, tạo bến xe dùng chung để hỗ trợ "
    "chuyển tuyến)."
)
make_table(
    ["Bảng (Table)", "Số lượng", "Mô tả dữ liệu"],
    [
        ["stops", "426 trạm",
         "60 trạm TP.HCM + 46 trạm Biên Hòa (dữ liệu mẫu biên soạn theo chuẩn GTFS) + "
         "320 trạm Kiên Giang (dữ liệu thật, 322 trạm gốc đã xử lý). Mỗi trạm: mã, tên "
         "(Việt + Anh), tọa độ (lat/lon), cờ bến trung chuyển, mã khu vực."],
        ["routes", "24 tuyến",
         "10 tuyến TP.HCM + 5 tuyến Biên Hòa + 9 tuyến Kiên Giang. Mỗi tuyến: số hiệu, "
         "lộ trình (Việt + Anh), giá vé phổ thông & sinh viên, giãn cách chạy (phút), "
         "khung giờ hoạt động."],
        ["route_stops", "455 dòng",
         "Quan hệ tuyến–trạm: thứ tự trạm trên tuyến (stop_sequence) và thời gian tích lũy "
         "từ đầu tuyến (offset_min) – dùng để tính thời gian di chuyển và mô phỏng vị trí xe."],
        ["search_logs", "Tăng dần theo lượt dùng",
         "Nhật ký tra cứu thực tế của người dùng: điểm đi, điểm đến, loại vé, số phương án "
         "tìm được, thời điểm. Chứng minh ứng dụng GHI dữ liệu lên Cloud Database."],
    ],
)

# ======================================================================
# 6. KET LUAN
# ======================================================================
h1("6. KẾT LUẬN & ĐỀ XUẤT PHÁT TRIỂN")
para(
    "Buổi thực hành số 4 về Deployment & Testing đã hoàn thành mục tiêu đề ra. Hệ thống Smart "
    "City Bus Assistant đã được đưa lên môi trường Internet với 4 thành phần điện toán đám mây "
    "hoạt động đồng bộ (Streamlit Community Cloud – Supabase PostgreSQL – Supabase PostgREST – "
    "Supabase Storage). Toàn bộ 12 Test Case đều cho kết quả PASS; các số liệu hiệu năng được đo "
    "thực tế. Ứng dụng có địa chỉ truy cập công khai, kết nối Cloud Database thật và ghi được dữ "
    "liệu lên Cloud."
)
para("Các hạn chế đã ghi nhận minh bạch:", bold=True)
bullet("Vị trí xe buýt trên bản đồ là ước tính mô phỏng theo biểu đồ chạy chuẩn – chưa tích hợp "
       "GPS thời gian thực (Việt Nam hiện chưa có API GPS xe buýt công khai miễn phí).")
bullet("Dữ liệu TP.HCM và Biên Hòa là bộ mẫu biên soạn theo cấu trúc GTFS (tọa độ trạm trung "
       "gian nội suy), không phải dữ liệu GTFS chính thức; riêng Kiên Giang là dữ liệu thật.")
bullet("Streamlit Community Cloud có cơ chế ngủ khi không truy cập lâu (khởi động lại ~30 giây).")

para("Định hướng nâng cấp tiếp theo:", bold=True)
bullet("Cài Cron-job ping định kỳ tới Web App để ngăn app rơi vào trạng thái ngủ, hoặc tự host "
       "trên VPS để chạy 24/7.")
bullet("Tích hợp API GPS thời gian thực khi có nguồn dữ liệu chính thức từ đơn vị vận hành.")
bullet("Mở rộng thêm nhiều tỉnh/thành; nâng cấp thuật toán tìm tuyến (nhiều lần chuyển tuyến, "
       "tối ưu đa tiêu chí thời gian – chi phí – số lần chuyển).")
bullet("Thêm chức năng \"Trạm gần tôi\" bằng định vị GPS của trình duyệt và tính năng lưu tuyến "
       "yêu thích (khi đó mới cần bổ sung đăng nhập).")
bullet("Đóng gói thành PWA / ứng dụng di động dùng chung Cloud API hiện tại.")

doc.save(OUT)
print("Da tao:", OUT)
